"""Datos ficticios. Ninguna prueba de este archivo contacta Supabase."""
from copy import deepcopy
from io import BytesIO
import json
from unittest.mock import Mock
from uuid import UUID, uuid4

from docx import Document
import pytest

from src.adapters.out.ai.generador_borrador_ollama import GeneradorBorradorOllama
from src.adapters.out.documents.exportador_docx import AVISO_PILOTO
from src.domain.entities.informe import Informe, VersionInforme
from src.domain.value_objects.estados import OrigenVersion
from tests.test_api import flujo_hasta_generar


ASUNTO = "PRUEBA NAXJI V2: solicitud de inspección del parque ficticio Aurora, sin resultados"


def preparar_v2(client):
    r = client.post("/solicitudes", json={"asunto": ASUNTO})
    assert r.status_code == 201, r.text
    sid = r.json()["id"]
    url = f"/solicitudes/{sid}"
    p = client.post(url + "/predecir-contexto").json()
    tipo = next(t for t in client.get("/tipos-informe").json() if t["codigo"] == "INFORME_TECNICO")
    area = client.get("/areas").json()[0]["id"]
    valida = client.post(url + "/validar-prediccion", json={"prediccion_id": p["id"], "resultado": "CORREGIDA",
                        "tipo_informe_id": tipo["id"], "area_destino_id": area, "normativa_ids": []})
    assert valida.status_code == 200, valida.text
    assert all(n["aceptada"] is False for n in client.get(url + "/contexto").json()["normativas"])
    plantilla = next(p for p in client.get("/plantillas", params={"tipo_informe_id": tipo["id"]}).json()
                     if p["version"] == 2 and "Piloto NAXJI" in p["nombre"])
    campos = client.get(f"/plantillas/{plantilla['id']}/campos").json()
    assert {c["clave"] for c in campos if c["obligatorio"]} == {"estado_resultados"}
    assert not {c["clave"] for c in campos} & {s["clave"] for s in plantilla["secciones_salida"]}
    valores = [{"campo_plantilla_id": c["id"], "valor": "Sin resultados de inspección"}
               for c in campos if c["clave"] == "estado_resultados"]
    body = {"asunto": ASUNTO, "tipo_informe_id": tipo["id"], "area_destino_id": area,
            "area_origen_id": area, "plantilla_id": plantilla["id"], "valores": valores}
    assert client.put(url + "/completa", json={**body, "valores": []}).status_code == 400
    guardada = client.put(url + "/completa", json=body)
    assert guardada.status_code == 200, guardada.text
    assert guardada.json()["estado"] == "LISTA_PARA_GENERAR"
    return sid, plantilla


def leer_docx(respuesta):
    assert respuesta.status_code == 200, respuesta.text if respuesta.status_code != 200 else ""
    assert respuesta.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    assert respuesta.headers["cache-control"] == "no-store"
    return Document(BytesIO(respuesta.content))


def test_fuentes_minimas_ollama_edicion_recuperacion_y_exportacion(client, deps):
    sid, plantilla = preparar_v2(client)
    propuesta = {s["clave"]: "Pendiente de verificación: no se aportaron resultados de inspección."
                 for s in plantilla["secciones_salida"]}
    llm = Mock(model="qwen2.5:7b", generar=Mock(return_value=json.dumps(propuesta)))
    deps.generar_borrador.generador = GeneradorBorradorOllama(llm)
    r = client.post(f"/solicitudes/{sid}/generar-borrador", json={})
    assert r.status_code == 201, r.text
    generado = r.json()
    entrada = json.loads(llm.generar.call_args.args[0])
    assert entrada["datos_ingresados"] == {"estado_resultados": "Sin resultados de inspección"}
    assert entrada["fuentes_sin_hechos"] is True
    assert entrada["contexto_confirmado"]["normas_seleccionadas_no_verificadas"] == []
    assert "no acredita hallazgos" in llm.generar.call_args.kwargs["sistema"]
    assert llm.generar.call_args.kwargs["esquema"]["properties"]["analisis_tecnico"]["pattern"] == "^Pendiente de verificación:"
    assert generado["modelo_ia"] == "qwen2.5:7b"
    assert generado["contenido"]["encabezado"]["asunto"] == ASUNTO
    iid = generado["informe_id"]
    contenido = {**generado["contenido"], "conclusiones": "Corrección humana ficticia: no se cuenta con resultados.\nPendiente de verificación presencial."}
    edicion = client.put(f"/informes/{iid}", json={"contenido": contenido, "numero_version": 1,
                         "titulo": "Informe de prueba corregido"})
    assert edicion.status_code == 200, edicion.text
    assert edicion.json()["numero_version"] == 2
    assert client.get(f"/solicitudes/{sid}/informe").json() == edicion.json()
    anterior = client.get(f"/informes/{iid}/versiones/1").json()
    assert anterior["contenido"] == generado["contenido"]
    assert anterior["titulo"] == "Informe Técnico" and anterior["titulo_versionado"] is True
    # Ninguna descarga puede llamar al proveedor ni alterar un registro.
    llm.generar.side_effect = AssertionError("No invocar Ollama al exportar")
    informe_antes = deepcopy(deps.informes.obtener_por_id(UUID(iid)))
    solicitud_antes = deepcopy(deps.solicitudes.obtener_por_id(UUID(sid)))
    for numero, contenido_esperado in [(1, generado["contenido"]), (2, contenido)]:
        r = client.get(f"/informes/{iid}/versiones/{numero}/docx")
        doc = leer_docx(r)
        assert f"-v{numero}.docx" in r.headers["content-disposition"]
        textos = [p.text for p in doc.paragraphs]
        assert textos[0] == ("Informe Técnico" if numero == 1 else "Informe de prueba corregido")
        assert AVISO_PILOTO in textos
        assert not doc.styles["Title"].element.xpath(".//w:pBdr")
        titulos = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        assert titulos == ["Encabezado"] + [s["titulo"] for s in plantilla["secciones_salida"]]
        for s in plantilla["secciones_salida"]:
            assert contenido_esperado[s["clave"]] in "\n".join(textos)
        assert "Asunto: " + ASUNTO in textos
        assert "[object Object]" not in "\n".join(textos)
    assert llm.generar.call_count == 1
    assert deps.informes.obtener_por_id(UUID(iid)) == informe_antes
    assert deps.solicitudes.obtener_por_id(UUID(sid)) == solicitud_antes


def test_sin_hechos_rechaza_salida_sin_advertencia_y_permite_reintentar(client, deps):
    sid, plantilla = preparar_v2(client)
    contenido = {s["clave"]: "Texto sin marcar la falta de evidencia." for s in plantilla["secciones_salida"]}
    llm = Mock(model="qwen2.5:7b", generar=Mock(return_value=json.dumps(contenido)))
    deps.generar_borrador.generador = GeneradorBorradorOllama(llm)
    r = client.post(f"/solicitudes/{sid}/generar-borrador", json={})
    assert r.status_code == 502 and r.json()["codigo"] == "OLLAMA_RESPUESTA_INVALIDA"
    assert client.get(f"/solicitudes/{sid}").json()["estado"] == "LISTA_PARA_GENERAR"
    assert not deps.memoria.informes
    llm.generar.return_value = json.dumps({k: "Pendiente de verificación: faltan resultados." for k in contenido})
    assert client.post(f"/solicitudes/{sid}/generar-borrador", json={}).status_code == 201


@pytest.mark.parametrize("token,status", [(None, 401), ("invalido", 401), ("demo-otro", 403),
                                          ("demo-revisor", 200), ("demo-admin", 200), ("demo-funcionario", 200)])
def test_descarga_aplica_permisos_del_informe(client, token, status):
    sid = flujo_hasta_generar(client)
    iid = client.post(f"/solicitudes/{sid}/generar-borrador", json={}).json()["informe_id"]
    client.headers.pop("Authorization")
    if token:
        client.headers["Authorization"] = "Bearer " + token
    for ruta in [f"/informes/{iid}/versiones/1/docx", f"/informes/{iid}/versiones/1", f"/solicitudes/{sid}/informe"]:
        r = client.get(ruta)
        assert r.status_code == status, r.text[:100]


def test_versiones_inexistentes_y_solicitudes_sin_informe(client):
    sid = flujo_hasta_generar(client)
    assert client.get(f"/solicitudes/{sid}/informe").status_code == 404
    iid = client.post(f"/solicitudes/{sid}/generar-borrador", json={}).json()["informe_id"]
    for numero, estado in [(0, 422), (-1, 422), (99, 404)]:
        assert client.get(f"/informes/{iid}/versiones/{numero}/docx").status_code == estado
    assert client.get(f"/informes/{uuid4()}/versiones/1/docx").status_code == 404


def test_exporta_legacy_y_texto_corregido_sin_consultar_plantilla_actual(client, deps):
    sid = flujo_hasta_generar(client)
    s = deps.solicitudes.obtener_por_id(UUID(sid))
    contenido = {"antecedentes": "A anterior", "desarrollo": "D anterior", "conclusiones": "C anterior",
                 "objetivo": "Texto antiguo adicional", "encabezado": {"asunto": "Asunto anterior", "area_origen": "Área ficticia"}}
    antiguo = Informe(s.id, s.plantilla_id, s.usuario_id, titulo="Informe anterior de prueba")
    antiguo.versiones.append(VersionInforme(antiguo.id, 1, contenido, OrigenVersion.IA, s.usuario_id))
    deps.informes.guardar(antiguo)
    deps.plantillas.plantillas[s.plantilla_id].activa = False
    r = client.put(f"/informes/{antiguo.id}", json={"numero_version": 1,
                   "contenido": {**contenido, "objetivo": "Objetivo antiguo corregido"}})
    assert r.status_code == 200, r.text
    doc = leer_docx(client.get(f"/informes/{antiguo.id}/versiones/2/docx"))
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 1"] == ["Encabezado", "Antecedentes", "Desarrollo", "Conclusiones", "Objetivo"]
    assert "Objetivo antiguo corregido" in [p.text for p in doc.paragraphs]
    assert client.get(f"/informes/{antiguo.id}/versiones/1").json()["contenido"] == contenido


def test_exportacion_caracteres_no_xml_error_controlado(client):
    sid = flujo_hasta_generar(client)
    generado = client.post(f"/solicitudes/{sid}/generar-borrador", json={}).json()
    iid = generado["informe_id"]
    assert client.put(f"/informes/{iid}", json={"numero_version": 1,
        "contenido": {**generado["contenido"], "conclusiones": "Texto\x00no XML"}}).status_code == 200
    assert client.get(f"/informes/{iid}/versiones/2/docx").status_code == 400
