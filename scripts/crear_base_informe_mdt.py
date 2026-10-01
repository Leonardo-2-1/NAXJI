"""Reconstrucción editable del Anexo 08 publicado como PDF; no es un DOCX oficial.

Autoridad: Directiva 001-2019-MDT/GM, páginas PDF 5, 7–10 y 19.
No copia sellos, firmas, siglas de 2019 ni el lema de una gestión anterior.
La identidad gráfica actual debe proporcionarla la municipalidad.
"""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def crear(destino):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = Cm(3.5)
    s.bottom_margin = s.left_margin = s.right_margin = Cm(3)
    s.header_distance = Cm(1)
    s.footer_distance = Cm(1.5)
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'es-PE')
    normal.element.get_or_add_rPr().append(lang)
    h = s.header.paragraphs[0]
    h.add_run('PILOTO NAXJI · Adaptación del Anexo 08 (2019), pendiente de validación\n').font.size = Pt(9)
    h.add_run('Logotipo institucional: [POR COMPLETAR]').font.size = Pt(9)
    h.paragraph_format.space_after = Pt(2)
    p = s.header.add_paragraph('Registro de documento: {{registro_documento}}\nExpediente: {{registro_expediente}}')
    p.alignment = 2
    for r in p.runs: r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(0)
    titulo = doc.add_paragraph()
    titulo.alignment = 1
    titulo.paragraph_format.space_after = Pt(14)
    titulo.paragraph_format.keep_with_next = True
    r = titulo.add_run('INFORME N° {{numero}}'); r.bold = r.underline = True
    for etiqueta, valor in [('A', '{{destinatario}}\n{{cargo_destinatario}}'),
                            ('ASUNTO', '{{asunto}}'), ('REF.', '{{referencia}}'),
                            ('FECHA', '{{lugar_emision}}, {{fecha}}')]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(3)
        p.paragraph_format.first_line_indent = Cm(-3)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3))
        p.paragraph_format.keep_with_next = True
        r = p.add_run(etiqueta + ':\t'); r.bold = True
        r = p.add_run(valor)
        if etiqueta in {'ASUNTO', 'REF.'}: r.bold = r.underline = True
        if etiqueta == 'FECHA':
            borders = OxmlElement('w:pBdr'); border = OxmlElement('w:bottom')
            for k,v in {'val':'single','sz':'6','space':'3','color':'000000'}.items(): border.set(qn('w:'+k),v)
            borders.append(border); p._p.get_or_add_pPr().append(borders)
            p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph('{{cuerpo}}')
    p.alignment = 3
    doc.add_paragraph('Atentamente,').paragraph_format.keep_with_next = True
    p = doc.add_paragraph('Firma y sello: [POR COMPLETAR]\n{{firmante}}\n{{cargo_firmante}}')
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.alignment = 2
    for valor, size in [('{{iniciales}}',9), ('Cc.: {{copias}}',10), ('Adjunto: {{anexos}}',10)]:
        p = doc.add_paragraph(valor)
        p.runs[0].font.size = Pt(size)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = not valor.startswith('Adjunto:')
    pie = s.footer.paragraphs[0]; pie.alignment = 2
    pie.paragraph_format.space_after = Pt(0)
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); pie._p.append(f)
    doc.core_properties.title = 'Base revisable Anexo 08 MDT 2019'
    doc.core_properties.author = 'NAXJI'
    doc.core_properties.comments = 'Adaptación PILOTO desde PDF. Pendiente de validación institucional y logotipo vigente. No es un formato oficial aprobado para NAXJI.'
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)


if __name__ == '__main__':
    crear(Path(__file__).resolve().parents[1] / 'src/adapters/out/documents/templates/mdt_anexo08_revision1.docx')
