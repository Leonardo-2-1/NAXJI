"""Word piloto: únicamente transforma datos persistidos en un documento editable."""
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.application.ports.output.exportador_informe import ExportadorInforme
from src.domain.services.errores import DatosInvalidos
from src.domain.value_objects.seccion_salida import CLAVES_TECNICAS, secciones_efectivas
from src.domain.value_objects.encabezado_documento import datos_oficiales, texto_publico, MARCADOR
from src.domain.value_objects.encabezado_documento import FORMATO_MDT
from .exportador_mdt import exportar_mdt


AVISO_PILOTO = "Formato Word piloto NAXJI. No es un formato institucional aprobado."
ETIQUETAS = {"asunto": "Asunto", "area_origen": "Área de origen", "area_destino": "Área de destino",
             "autor_id": "ID del autor", "tipo_informe_id": "ID del tipo de informe",
             "area_destino_id": "ID del área de destino", "normativa_ids": "IDs de normas confirmadas"}


def titulo_dato(clave):
    return ETIQUETAS.get(clave, clave.replace("_", " ").capitalize())


def secciones_documento(version):
    secciones = secciones_efectivas(version.secciones_salida)
    if version.secciones_salida is None:
        conocidas = {s.clave for s in secciones} | CLAVES_TECNICAS
        # Mismo tratamiento que el editor para textos anteriores al paso 3.
        secciones += [SimpleNamespace(clave=k, titulo=titulo_dato(k)) for k, v in version.contenido.items()
                      if k not in conocidas and isinstance(v, str)]
    return secciones


class ExportadorDocx(ExportadorInforme):
    def exportar(self, informe, version):
        encabezado = version.contenido.get("encabezado", {})
        formato = encabezado.get("formato_documento") if isinstance(encabezado, dict) else None
        if formato == FORMATO_MDT:
            return exportar_mdt(version)
        if formato:
            raise DatosInvalidos("El formato documental guardado no está disponible")
        documento = Document()
        pagina = documento.sections[0]
        pagina.page_width, pagina.page_height = Cm(21), Cm(29.7)
        pagina.top_margin = pagina.bottom_margin = Cm(1.8)
        pagina.footer_distance = Cm(.9)
        pagina.left_margin = pagina.right_margin = Cm(2.5)
        # La plantilla base de python-docx incluye un borde azul en Title.
        # Este formato piloto usa solo jerarquía tipográfica, sin líneas decorativas.
        for borde in documento.styles.element.xpath(".//w:pBdr"):
            borde.getparent().remove(borde)
        for nombre, tamano in [("Normal", 11), ("Title", 16), ("Heading 1", 12), ("Heading 2", 11)]:
            estilo = documento.styles[nombre]
            estilo.font.name = "Arial"
            estilo.font.size = Pt(tamano)
            estilo.font.color.rgb = RGBColor(0, 0, 0)
            estilo.paragraph_format.space_after = Pt(5)
            if nombre != "Normal":
                estilo.paragraph_format.keep_with_next = True
                estilo.paragraph_format.space_before = Pt(8 if nombre.startswith("Heading") else 0)
        normal = documento.styles["Normal"]
        normal.paragraph_format.line_spacing = 1.05
        normal.paragraph_format.widow_control = True
        idioma = OxmlElement("w:lang")
        idioma.set(qn("w:val"), "es-PE")
        normal.element.get_or_add_rPr().append(idioma)
        documento.core_properties.author = "NAXJI"
        documento.core_properties.last_modified_by = "NAXJI"
        documento.core_properties.comments = ""
        pie = pagina.footer.paragraphs[0]
        pie.alignment = 2
        pie.add_run("NAXJI · PILOTO · Página ").font.size = Pt(9)
        for indice, campo in enumerate(("PAGE", "NUMPAGES")):
            if indice:
                pie.add_run(" de ").font.size = Pt(9)
            nodo = OxmlElement("w:fldSimple")
            nodo.set(qn("w:instr"), campo)
            pie._p.append(nodo)
        titulo = version.titulo if version.titulo is not None else informe.titulo
        try:
            documento.core_properties.title = titulo or "Informe"
            documento.add_paragraph(titulo or "Informe", "Title")
            aviso = documento.add_paragraph(AVISO_PILOTO)
            aviso.paragraph_format.keep_with_next = True
            aviso.runs[0].font.size = Pt(9)
            h = version.contenido.get("encabezado")
            h = h if isinstance(h, dict) else {}
            origen = texto_publico(h.get("plantilla_nombre")) or "Plantilla histórica"
            revision = h.get("plantilla_version")
            revision = str(revision) if isinstance(revision, int) else "no registrada"
            detalle = documento.add_paragraph(f"{origen} · Versión de plantilla {revision}\nVersión guardada {version.numero_version} del borrador · Sujeto a revisión humana")
            detalle.runs[0].font.size = Pt(9)
            detalle.paragraph_format.keep_with_next = True
            if version.titulo is None:
                documento.add_paragraph("Versión anterior sin título propio guardado. Se muestra el título actual del informe.")
            documento.add_heading("Encabezado", level=1)
            oficiales = datos_oficiales(h)
            # Lista explícita: nunca exporta autor_id ni UUID técnicos de versiones antiguas.
            filas = [("Número", oficiales["numero"]), ("A", oficiales["destinatario"]),
                     ("De", oficiales["emisor"]), ("Asunto", texto_publico(h.get("asunto"))),
                     ("Referencia", oficiales["referencia"]), ("Fecha del documento", oficiales["fecha"])]
            for etiqueta, valor in filas:
                p = documento.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.keep_with_next = True
                p.add_run(etiqueta + ": ").bold = True
                p.add_run(valor or MARCADOR)
            grupos = []
            for seccion in secciones_documento(version):
                grupo = [documento.add_heading(seccion.titulo, level=1)]
                # No reemplaza un vacío por hechos nuevos. Incluye también opcionales sin texto.
                texto = version.contenido.get(seccion.clave, "")
                if not isinstance(texto, str):
                    raise DatosInvalidos("La versión contiene una sección que no es texto")
                for linea in texto.split("\n"):
                    p = documento.add_paragraph(linea)
                    p.paragraph_format.keep_together = len(linea) < 900
                    grupo.append(p)
                grupos.append(grupo)
            # Una sección breve final viaja con el cierre anterior, no queda sola.
            # Se acota por longitud para no encadenar bloques largos a otra página.
            if len(grupos) >= 2 and sum(len(p.text) for g in grupos[-2:] for p in g) < 1400:
                cierre = [p for g in grupos[-2:] for p in g]
                for p in cierre[:-1]:
                    p.paragraph_format.keep_with_next = True
            elif len(grupos) >= 2 and sum(len(p.text) for p in grupos[-1]) < 900:
                # Conclusión extensa: conserva al menos su cierre junto a la sección final.
                grupos[-2][-1].paragraph_format.keep_with_next = True
                for p in grupos[-1][:-1]:
                    p.paragraph_format.keep_with_next = True
            firma = documento.add_paragraph()
            firma.paragraph_format.keep_together = True
            firma.paragraph_format.space_before = Pt(10)
            firma.add_run("Firmante: ").bold = True
            firma.add_run(oficiales["firmante"] or MARCADOR)
            firma.add_run("\nCargo: ").bold = True
            firma.add_run(oficiales["cargo_firmante"] or MARCADOR)
            if grupos:
                grupos[-1][-1].paragraph_format.keep_with_next = True
            salida = BytesIO()
            documento.save(salida)
            return salida.getvalue()
        except ValueError:
            raise DatosInvalidos("La versión contiene caracteres no compatibles con Word; revise el texto antes de descargar") from None
