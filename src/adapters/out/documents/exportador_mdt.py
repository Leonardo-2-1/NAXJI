"""Llena la base revisable del Anexo 08 con la versión guardada, nunca con IA."""
from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.text.paragraph import Paragraph

from src.domain.services.errores import DatosInvalidos
from src.domain.value_objects.encabezado_documento import MARCADOR, datos_oficiales, texto_publico

BASE = Path(__file__).parent / 'templates/mdt_anexo08_revision1.docx'


def exportar_mdt(version):
    if [s.clave for s in version.secciones_salida or []] != ['cuerpo']:
        raise DatosInvalidos('La versión guardada no coincide con la estructura del Anexo 08')
    cuerpo = version.contenido.get('cuerpo')
    if not isinstance(cuerpo, str) or not cuerpo.strip():
        raise DatosInvalidos('Falta el cuerpo del informe interno guardado')
    h = version.contenido.get('encabezado', {})
    datos = {**datos_oficiales(h), 'asunto': texto_publico(h.get('asunto'))}
    if datos.get('fecha'):
        try:
            fecha = date.fromisoformat(datos['fecha'])
            meses = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio',
                     'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')
            datos['fecha'] = f'{fecha.day:02d} de {meses[fecha.month-1]} de {fecha.year}'
        except ValueError:
            raise DatosInvalidos('La fecha administrativa guardada no es válida') from None
    doc = Document(BASE)
    doc.core_properties.title = version.titulo or 'Informe interno de inspección'
    doc.core_properties.last_modified_by = 'NAXJI'
    try:
        for p in [*doc.paragraphs, *doc.sections[0].header.paragraphs]:
            if p.text == '{{cuerpo}}':
                continue
            for r in p.runs:
                # Sustitución de una sola pasada sobre la base; ni los valores ni
                # el cuerpo del usuario se reinterpretan como variables.
                r.text = re.sub(r'\{\{([a-z_]+)\}\}', lambda m: datos.get(m[1]) or MARCADOR, r.text)
        for p in list(doc.paragraphs):
            if p.text == '{{cuerpo}}':
                lineas = cuerpo.split('\n')
                # Conservar párrafos y su estilo de la base sin interpretar HTML/Markdown.
                for i, linea in enumerate(lineas):
                    nuevo = deepcopy(p._p)
                    for r in nuevo.xpath('./w:r'): nuevo.remove(r)
                    parrafo = Paragraph(nuevo, p._parent)
                    parrafo.add_run(linea)
                    parrafo.paragraph_format.keep_together = len(linea) < 700
                    parrafo.paragraph_format.keep_with_next = i == len(lineas)-1
                    p._p.addprevious(nuevo)
                p._p.getparent().remove(p._p)
        salida = BytesIO(); doc.save(salida)
        return salida.getvalue()
    except ValueError:
        raise DatosInvalidos('La versión contiene caracteres no compatibles con Word; revise el texto antes de descargar') from None
